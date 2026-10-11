from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(r"C:\Users\khorh\smart-wardrobe")
TEMPLATE = Path(r"C:\Users\khorh\Downloads\Paper.docx")
OUT = ROOT / "deliverables" / "Khor_Hua_Sheng_Smart_Wardrobe_Summary_Paper_Eight_Part.docx"


def set_font(run, size=10.5, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, colour="D9D9D9"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        elem = borders.find(tag)
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:color"), colour)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def remove_all_body_content(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def add_text(doc, text, *, bold=False, italic=False, size=10.5, alignment=None, before=0, after=6, first_indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(first_indent)
    if alignment is not None:
        p.alignment = alignment
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        set_font(r, size=11, bold=True)
    else:
        p.paragraph_format.left_indent = Inches(0.18)
        r = p.add_run(text)
        set_font(r, size=10.5, bold=True)
    return p


def add_table(doc, caption, headers, rows, widths, page_break_before=False):
    cap = add_text(doc, caption, italic=True, size=9.5, before=5, after=3)
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.page_break_before = page_break_before
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, label in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        set_cell_shading(cell, "1F4E79")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_font(r, size=8.7, bold=True)
        r.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for row_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_i % 2 == 1:
                set_cell_shading(cells[i], "F2F6FA")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(value)
            set_font(r, size=8.4)
        prevent_row_split(table.rows[-1])
    return table


def add_architecture_figure(doc):
    caption = add_text(doc, "Figure 1. Simplified architecture for the garment-analysis and outfit-visualisation subsystem.", italic=True, size=9.5, before=5, after=3)
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.page_break_before = True
    table = doc.add_table(rows=4, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, colour="B7C9D6")
    layers = [
        ("Presentation tier", "Mobile-responsive web frontend", "DDEBF7"),
        ("Application tier", "FastAPI modular monolith: validated processing jobs, candidate review, transactional persistence and outfit-visualisation orchestration", "EAF2F8"),
        ("AI integration boundaries", "Segmentation; visible attributes and embeddings; VTOFF reconstruction; VTON generation", "F5F9FC"),
        ("Data tier", "PostgreSQL with pgvector and protected image/file storage", "DDEBF7"),
    ]
    for row, (label, detail, fill) in zip(table.rows, layers):
        cell = row.cells[0]
        cell.width = Inches(6.2)
        set_cell_shading(cell, fill)
        set_cell_margins(cell, top=110, start=150, bottom=110, end=150)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label)
        set_font(r, size=9.2, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(detail)
        set_font(r2, size=8.5)
        prevent_row_split(row)
    return table


doc = Document(TEMPLATE)
remove_all_body_content(doc)
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

title = add_text(doc, "Development and Evaluation of a Deep Learning Based Garment Analysis and Outfit Visualisation System", bold=True, size=14, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=2)
title.paragraph_format.keep_with_next = True
author = add_text(doc, "Khor Hua Sheng", size=10.5, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=10)
author.paragraph_format.keep_with_next = True

add_heading(doc, "ABSTRACT", 1)
add_text(doc, "Personal wardrobe photographs are difficult to convert into reliable digital records because one image may contain several garments, background clutter and uncertain visual attributes. This study develops and evaluates the garment-analysis and outfit-visualisation subsystem of a Smart Wardrobe mobile web application. The subsystem separates a validated wardrobe photograph into reviewable garment candidates, stores only user-confirmed records, and produces an explicitly labelled AI-generated appearance preview for supported outfits. Iterative development integrated instance segmentation, optional virtual try-off reconstruction, visible attribute suggestions, visual embeddings, protected media storage and virtual try-on through model-independent application interfaces. Evaluation combined requirement-based functional audits, held-out segmentation metrics and paired reconstruction metrics. Held-out DeepFashion2 evaluation produced box mAP50 of 0.753 and mask mAP50 of 0.732. In a fixed 30-pair DressCode audit, virtual try-off reconstruction achieved mean SSIM of 0.855, LPIPS of 0.171 and DISTS of 0.188; the supported virtual try-on audit achieved 0.912, 0.072 and 0.076 respectively. These image-similarity metrics do not measure physical fit. The results show a traceable workflow from upload to confirmed wardrobe records while retaining user control over uncertain model outputs.", after=5, first_indent=0.25)
add_text(doc, "Keywords: digital wardrobe, garment segmentation, virtual try-off, virtual try-on, computer vision", italic=True, size=10, after=10)

add_heading(doc, "1. INTRODUCTION", 1)
add_heading(doc, "1.1 Background", 2)
add_text(doc, "Personal wardrobes are often represented by unstructured photographs or scattered physical items. A useful digital wardrobe must turn those photographs into separate, searchable garment records without treating uncertain computer-vision output as established fact.", first_indent=0.25)
add_heading(doc, "1.2 Problem Statement", 2)
add_text(doc, "A single wardrobe image can contain several overlapping garments, inconsistent backgrounds and incomplete visual evidence. Manual entry is time-consuming, whereas automatic extraction can assign an incorrect category, colour or attribute. Generated garment and try-on images can also alter visible details.", first_indent=0.25)
add_heading(doc, "1.3 Research Gap", 2)
add_text(doc, "DeepFashion2 supports instance-level clothing analysis (Ge et al., 2019), and Fashionpedia provides a vocabulary for localised visible attributes (Jia et al., 2020). Virtual try-on studies demonstrate plausible image synthesis, but these capabilities are often evaluated separately from the user-controlled creation of private wardrobe records. Therefore, this study proposes the Smart Wardrobe mobile web application, whose subsystem connects extraction, human verification, storage and outfit visualisation with explicit evidence boundaries.", first_indent=0.25)
add_heading(doc, "1.4 Objectives", 2)
add_text(doc, "This subsystem aims to: (i) process one wardrobe image into separate garment candidates; (ii) create editable, evidence-linked garment representations for downstream use; and (iii) generate a whole-outfit appearance preview only from confirmed garments and an authorised person image.", first_indent=0.25)
add_heading(doc, "1.5 Contributions", 2)
add_text(doc, "The contribution is a working, model-aware subsystem that retains segmentation evidence, requires user confirmation before persistence, stores visual embeddings for later retrieval, and labels generated media as an appearance preview rather than a prediction of size, comfort or physical fit.", first_indent=0.25)

add_heading(doc, "2. LITERATURE REVIEW", 1)
add_heading(doc, "2.1 Related Applications", 2)
add_text(doc, "Digital wardrobe applications commonly catalogue garments and support outfit planning. Their value depends on accurate garment records, but conventional workflows frequently require a user to crop images and enter descriptions manually. The present subsystem focuses on reducing that effort while keeping the user in control of the final record.", first_indent=0.25)
add_heading(doc, "2.2 Existing Methods", 2)
add_text(doc, "Instance segmentation separates individual clothing items from an image. Fashionpedia's Attribute-Mask R-CNN is an established joint segmentation and attribute-prediction baseline (Jia et al., 2020). VITON and later high-resolution virtual try-on methods synthesise a target appearance from person and garment images (Han et al., 2018; Lee et al., 2022). These methods are useful building blocks, but their visual outputs remain probabilistic.", first_indent=0.25)
add_heading(doc, "2.3 Comparison of Existing Solutions", 2)
add_text(doc, "An end-to-end attribute model can directly associate labels with detected instances, whereas a modular pipeline can independently update segmentation, attribute inference and image synthesis. The selected design uses segmentation to retain observed garment crops, then adds editable visible-attribute suggestions and optional generated alternatives. This makes the original evidence available when an inferred or generated result is unsuitable.", first_indent=0.25)
add_heading(doc, "2.4 Research Gap", 2)
add_text(doc, "Existing benchmarks establish model-level performance, but they do not by themselves provide an auditable private-wardrobe workflow. This study addresses that gap by combining model adapters with ownership checks, review decisions, atomic persistence, failure states and provenance for AI-generated media.", first_indent=0.25)

add_heading(doc, "3. METHODOLOGY", 1)
add_heading(doc, "3.1 Development Methodology", 2)
add_text(doc, "An iterative and incremental approach was selected because the workflow and AI components required repeated refinement. Each increment preserved the application contract while model adapters isolated segmentation, virtual try-off, visible attribute inference and virtual try-on behaviour.", first_indent=0.25)
add_heading(doc, "3.2 Requirements Analysis", 2)
add_text(doc, "The primary user is a wardrobe owner. The core functional requirements are validated image entry, multi-garment candidate extraction, individual acceptance or rejection, editable controlled metadata, protected storage, retrieval-ready visual embeddings and authorised outfit visualisation. Non-functional requirements include visible progress, cancellation, failure-safe persistence, owner scoping and traceable media provenance.", first_indent=0.25)
add_heading(doc, "3.3 System Design", 2)
add_text(doc, "The design consists of a mobile web interface, FastAPI application services, PostgreSQL-compatible persistence and protected media storage. A processing job normalises an uploaded JPEG or PNG and invokes instance segmentation. Each candidate retains a source-space mask, bounding box, category, confidence and segmentation-guided crop. Accepted candidates are committed transactionally as garments or as members of one outfit.", first_indent=0.25)
add_architecture_figure(doc)
add_heading(doc, "3.4 Application Development", 2)
add_text(doc, "The mobile-responsive web frontend was selected to provide one user-facing interface, while FastAPI centralises authentication, validation, transactions and access to the controlled Python and GPU environment required by the AI runtime. PostgreSQL with pgvector keeps reviewed records and vector representations in the same transactional data environment, while protected file storage separates private image media from relational records. Model-independent adapters prevent model-specific libraries from being embedded in application services. A frozen SigLIP visual encoder with project-trained Fashionpedia attribute heads supplies editable visible labels and a 768-dimensional embedding. Optional TryOffDiff reconstruction is presented separately from the observed crop. For supported outfits, a pinned CatVTON runtime applies reviewed garments sequentially in layer order.", first_indent=0.25)
add_heading(doc, "3.5 Testing and Evaluation", 2)
add_text(doc, "Testing combined functional acceptance audits, a held-out DeepFashion2 segmentation conversion evaluation, blind checks of selected visual fields and paired image-similarity audits. The reconstruction audits used a fixed-seed, stratified 30-pair sample from authorised DressCode test manifests: ten upper-body, ten lower-body and ten dress examples. SSIM, LPIPS and DISTS were reported with qualitative failure labels. These measures compare images; they do not validate physical fit.", first_indent=0.25)

add_heading(doc, "4. RESULTS", 1)
add_heading(doc, "4.1 Application Features", 2)
add_text(doc, "The completed subsystem provides upload progress and cancellation, per-candidate review, controlled metadata selections, retention of the original segmentation crop, optional VTOFF comparison, visual embedding creation and whole-outfit visualisation. The user can correct predictions before saving, and only accepted candidates enter the wardrobe.", first_indent=0.25)
add_heading(doc, "4.2 Functional Testing", 2)
add_text(doc, "The requirement audits test implemented behaviour and evidence rather than model accuracy alone. Table 1 states what each mapped group covers so that the pass counts are interpretable.", first_indent=0.25)
add_table(doc, "Table 1. Requirement-audit scope and accepted behaviour.", ["Audit", "Requirement scope", "What passed"], [
    ["O1: 25/25", "FR-01-05; FR-10-24; FR-37-39; AIR-01-02", "Image entry and validation; instance segmentation; candidate review; atomic garment/outfit saving; owner-scoped operations; editing and safe archive; cancellation integrity; dataset conversion and held-out segmentation evidence."],
    ["O2: 12 groups", "FR-06-09; grouped FR-10-18, FR-20-21 and FR-38; AIR-03-05, AIR-07", "Optional VTOFF with crop fallback; structured attributes and 768-D embeddings; review/persistence/ownership boundaries; paired VTOFF metrics; generated-media review; attribute and embedding failure evidence; reproducible experiment records."],
    ["O3: 18/18", "FR-18; FR-25-39; AIR-06-07", "Visualisation entry points; source and person-image validation; protected output; lifecycle states; AI and physical-fit notice; retry, fail-closed completion, cancellation and provenance; VTON runtime evidence."],
], [0.8, 1.8, 3.65], page_break_before=True)
add_heading(doc, "4.3 User Evaluation", 2)
add_text(doc, "A formal usability study with wardrobe owners has not yet been conducted. Therefore, this paper does not claim user satisfaction, usability scores or recommendation quality. The current evaluation is limited to functional, segmentation, representation and paired image-reconstruction evidence.", first_indent=0.25)
add_heading(doc, "4.4 Results", 2)
add_text(doc, "Table 2 summarises the measured results. Weak colour and dominance predictions remain editable suggestions rather than automatic facts.", first_indent=0.25)
add_table(doc, "Table 2. Summary of held-out and paired evaluation evidence.", ["Evaluation", "Result", "Interpretation"], [
    ["Held-out segmentation", "Box mAP50 0.753; mask mAP50 0.732", "Provides candidate evidence, but user review remains necessary."],
    ["Visible attributes", "Colour temperature accuracy 89.7%; primary colour 35.9%; dominance 38.5%", "Only the predictions are prefilled; the user may modify them before saving."],
    ["VTOFF reconstruction, n = 30", "SSIM 0.855; LPIPS 0.171; DISTS 0.188", "The generated reconstruction exceeded the retained crop baseline in the paired audit."],
    ["VTON reconstruction, n = 30", "SSIM 0.912; LPIPS 0.072; DISTS 0.076", "Bounded paired image-similarity evidence across upper, lower and dress groups."],
    ["Visual embeddings", "Top-1 category retrieval 47.1%; macro F1 46.6%; MRR 0.601", "Available for downstream retrieval; not evidence of recommender quality."],
], [1.35, 1.85, 3.05])

add_heading(doc, "5. DISCUSSION", 1)
add_heading(doc, "5.1 Findings", 2)
add_text(doc, "The results support the decision to make AI output assistive rather than authoritative. Segmentation creates the observed evidence used for a candidate record, while the user accepts the garment, corrects descriptive fields and chooses between the observed crop and an optional generated reconstruction. The low primary-colour and dominance results illustrate why automated metadata must remain reviewable.", first_indent=0.25)
add_heading(doc, "5.2 Comparison with Previous Studies", 2)
add_text(doc, "The implementation uses the same broad problem domains addressed by DeepFashion2, Fashionpedia and virtual try-on studies, but it is not presented as a direct performance replacement for their specialist benchmarks. Its distinctive focus is the integration of extraction, controlled review, persistence and outfit-level visualisation in one private-wardrobe workflow. Fashionpedia's official Attribute-Mask R-CNN remains an appropriate future benchmark for the attribute module.", first_indent=0.25)
add_heading(doc, "5.3 Implications", 2)
add_text(doc, "For the wider Smart Wardrobe system, the confirmed category, controlled visible attributes and 768-dimensional embedding give the recommender subsystem a reviewable garment representation rather than unverified raw predictions. The provenance and disclaimer controls also make clear that a generated outfit image is a visual aid, not a fitting decision.", first_indent=0.25)
add_heading(doc, "5.4 Limitations", 2)
add_text(doc, "The 30-pair reconstruction audits are development-size samples, not full public benchmarks or user studies. The image metrics cannot establish size, comfort, fabric behaviour or real-world fit. Some difficult examples altered colour, print, neckline, sleeve, hem or drape. Multi-garment processing is supported sequentially, but no paired ground-truth evaluation yet covers complete multi-garment outfits.", first_indent=0.25)

add_heading(doc, "6. CONCLUSION", 1)
add_heading(doc, "6.1 Summary", 2)
add_text(doc, "This study developed and evaluated the garment-analysis and outfit-visualisation subsystem of a Smart Wardrobe application. It transforms a wardrobe photograph into separate reviewable candidates, saves only user-confirmed garments and outfits, and generates traceable AI appearance previews from authorised media. The stated evidence shows that the workflow operates within its defined scope; it does not claim perfect fashion understanding or physical-fit prediction.", first_indent=0.25)
add_heading(doc, "6.2 Future Work", 2)
add_text(doc, "Future work should compare the project-trained attribute heads with Fashionpedia's official Attribute-Mask R-CNN baseline on a matched held-out set, expand evaluation of whole-outfit outputs, improve difficult clothing categories, and conduct a structured usability study with representative wardrobe owners. The recommender subsystem should subsequently be evaluated separately using the confirmed garment records.", first_indent=0.25)

add_heading(doc, "REFERENCES", 1)
references = [
    "Ge, Y., Zhang, R., Wang, X., Tang, X., & Luo, P. (2019). DeepFashion2: A versatile benchmark for detection, pose estimation, segmentation and re-identification of clothing images. Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition, 5337-5345. https://doi.org/10.1109/CVPR.2019.00548",
    "Han, X., Wu, Z., Wu, Z., Yu, R., & Davis, L. S. (2018). VITON: An image-based virtual try-on network. Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition, 7543-7552. https://doi.org/10.1109/CVPR.2018.00787",
    "Jia, M., Shi, M., Sirotenko, M., Cui, Y., Cardie, C., Hariharan, B., Adam, H., & Belongie, S. (2020). Fashionpedia: Ontology, segmentation, and an attribute localization dataset. Computer Vision - ECCV 2020, 316-332. https://doi.org/10.1007/978-3-030-58452-8_19",
    "Lee, S., Gu, G., Park, S., Choi, S., & Choo, J. (2022). High-resolution virtual try-on with misalignment and occlusion-handled conditions. Computer Vision - ECCV 2022. https://doi.org/10.1007/978-3-031-19790-1_13",
    "Velioglu, R., Bevandic, P., Chan, R., & Hammer, B. (2025a). TryOffDiff: Virtual-Try-Off via high-fidelity garment reconstruction using diffusion models. Proceedings of the British Machine Vision Conference.",
    "Velioglu, R., Bevandic, P., Chan, R., & Hammer, B. (2025b). MGT: Extending virtual try-off to multi-garment scenarios. Proceedings of the IEEE/CVF International Conference on Computer Vision Workshops.",
]
for i, ref in enumerate(references, 1):
    p = add_text(doc, f"{i}. {ref}", size=9.5, after=3)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
